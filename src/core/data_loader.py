import pandas as pd
import numpy as np
from pathlib import Path
from typing import Union, Optional, Dict, Any, List
import json, pickle, io, os, re
from PIL import Image
import pytesseract
import fitz
import docx
import textract
from bs4 import BeautifulSoup
import yaml
import csv
import requests
from io import BytesIO
import base64

class UniversalDataLoader:
    SUPPORTED_EXTENSIONS = {
        '.csv':'tabular','.tsv':'tabular','.xlsx':'tabular','.xls':'tabular',
        '.parquet':'tabular','.feather':'tabular','.json':'tabular','.jsonl':'tabular',
        '.xml':'tabular','.html':'tabular','.txt':'text','.md':'text',
        '.pdf':'text','.docx':'text','.doc':'text','.rtf':'text','.odt':'text',
        '.jpg':'image','.jpeg':'image','.png':'image','.gif':'image',
        '.bmp':'image','.tiff':'image','.webp':'image','.svg':'image',
        '.wav':'audio','.mp3':'audio','.flac':'audio','.ogg':'audio','.m4a':'audio',
        '.mp4':'video','.avi':'video','.mov':'video','.mkv':'video',
        '.pkl':'ml_model','.joblib':'ml_model','.h5':'ml_model','.pt':'ml_model',
        '.pth':'ml_model','.onnx':'ml_model','.tflite':'ml_model',
        '.log':'text','.sql':'text','.py':'text','.ipynb':'notebook','.pickle':'binary',
    }
    def __init__(self):
        self.data_type = None
        self.metadata = {}
    def load(self, file_path: Union[str, Path], **kwargs) -> Dict[str, Any]:
        file_path = Path(file_path)
        ext = file_path.suffix.lower()
        data_type = self.SUPPORTED_EXTENSIONS.get(ext, self._detect_data_type(file_path))
        self.data_type = data_type
        loaders = {
            'tabular': self._load_tabular,
            'text': self._load_text,
            'image': self._load_image,
            'audio': self._load_audio,
            'video': self._load_video,
            'ml_model': self._load_ml_model,
            'notebook': self._load_notebook,
            'binary': self._load_binary,
        }
        if data_type in loaders:
            data = loaders[data_type](file_path, **kwargs)
            # If the loaded data is not a DataFrame but it's a dict/list, treat as JSON
            if data_type == 'tabular' and not isinstance(data, pd.DataFrame):
                data_type = 'json'
                self.data_type = 'json'
        else:
            data = self._load_generic(file_path)
        return {'data':data, 'type':data_type, 'metadata':self.metadata,
                'original_file':str(file_path), 'file_size':file_path.stat().st_size}
    def _detect_data_type(self, file_path):
        with open(file_path, 'rb') as f:
            header = f.read(1024)
        if b'\0' in header and not self._is_text_based(header):
            if header.startswith(b'\x89PNG') or header.startswith(b'\xff\xd8'): return 'image'
            if header.startswith(b'PK') and b'[Content_Types]' in header: return 'text'
            if header.startswith(b'%PDF'): return 'text'
            return 'binary'
        try:
            content = header.decode('utf-8')
            if '<' in content and '>' in content: return 'tabular' if 'table' in content else 'text'
            if '{' in content and '}' in content: return 'tabular'
            if ',' in content or '\t' in content: return 'tabular'
            return 'text'
        except:
            return 'binary'
    def _is_text_based(self, header):
        try: header.decode('utf-8'); return True
        except: return False
    def _load_tabular(self, file_path, **kwargs):
        ext = file_path.suffix.lower()
        strategies = []
        if ext == '.csv':
            strategies = [
                lambda: pd.read_csv(file_path, **kwargs),
                lambda: pd.read_csv(file_path, encoding='latin1', **kwargs),
                lambda: pd.read_csv(file_path, sep=';', **kwargs),
                lambda: pd.read_csv(file_path, sep='\t', **kwargs),
            ]
        elif ext == '.json' or ext == '.jsonl':
            # Try multiple JSON strategies, fallback to raw JSON/JSONL
            strategies = [
                lambda: pd.read_json(file_path, **kwargs),
                lambda: pd.read_json(file_path, lines=True, **kwargs),
                lambda: pd.json_normalize(json.load(open(file_path))),
            ]
        elif ext in ['.xlsx', '.xls']:
            strategies = [
                lambda: pd.read_excel(file_path, **kwargs),
                lambda: pd.read_excel(file_path, engine='openpyxl', **kwargs),
            ]
        else:
            strategies = [lambda: pd.read_csv(file_path, **kwargs)]

        # Try each strategy
        for strategy in strategies:
            try:
                df = strategy()
                if len(df) > 0:
                    self.metadata['rows'] = len(df)
                    self.metadata['columns'] = len(df.columns)
                    return df
            except:
                continue

        # ----- FALLBACK for JSON / JSONL -----
        if ext == '.json' or ext == '.jsonl':
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read().strip()
                # Try single JSON object
                try:
                    raw_data = json.loads(content)
                except json.JSONDecodeError:
                    # Try line‑delimited JSON (JSONL)
                    lines = content.splitlines()
                    parsed_objects = []
                    for line in lines:
                        line = line.strip()
                        if line:
                            try:
                                parsed_objects.append(json.loads(line))
                            except json.JSONDecodeError:
                                # skip malformed lines
                                pass
                    if parsed_objects:
                        raw_data = parsed_objects
                    else:
                        raise ValueError("Could not parse any JSON objects from file")
                self.metadata['json_type'] = type(raw_data).__name__
                self.metadata['is_tabular'] = False
                return raw_data
            except Exception as e:
                raise ValueError(f"Could not read JSON file: {e}")

        # If all else fails, raise an informative error
        raise ValueError(f"Could not load {file_path} as tabular. "
                        f"If this is a JSON file, ensure it's a valid JSON array or line‑delimited JSON.")
    def _load_text(self, file_path, **kwargs):
        ext = file_path.suffix.lower()
        text_content = ""
        structured_data = {}
        try:
            if ext == '.pdf':
                text_content, structured_data = self._extract_pdf(file_path)
            elif ext in ['.docx','.doc']:
                text_content = self._extract_docx(file_path)
            elif ext == '.txt':
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    text_content = f.read()
            else:
                text_content = textract.process(file_path).decode('utf-8')
            self.metadata['word_count'] = len(text_content.split())
            self.metadata['char_count'] = len(text_content)
            self.metadata['sentences'] = len(text_content.split('.'))
            return {'raw_text':text_content, 'structured_data':structured_data,
                    'metadata':self.metadata, 'file_type':ext}
        except Exception as e:
            raise ValueError(f"Error loading text: {e}")
    def _extract_pdf(self, file_path):
        doc = fitz.open(file_path)
        text_content = ""
        tables = []
        for page_num, page in enumerate(doc):
            text_content += f"\n--- Page {page_num+1} ---\n{page.get_text()}"
            t = page.find_tables()
            if t.tables:
                for table in t.tables:
                    rows = [ [str(c).strip() for c in row] for row in table.extract() ]
                    if rows: tables.append({'page':page_num+1, 'data':rows,
                                            'row_count':len(rows), 'col_count':len(rows[0])})
        doc.close()
        self.metadata['pages'] = len(doc)
        return text_content, {'tables':tables}
    def _extract_docx(self, file_path):
        doc = docx.Document(file_path)
        return '\n'.join([para.text for para in doc.paragraphs])
    def _load_image(self, file_path, **kwargs):
        img = Image.open(file_path)
        self.metadata['format']=img.format; self.metadata['size']=img.size; self.metadata['mode']=img.mode
        try:
            ocr_text = pytesseract.image_to_string(img)
            self.metadata['ocr_text'] = ocr_text
        except:
            self.metadata['ocr_text'] = ""
        return {'image':img, 'array':np.array(img), 'metadata':self.metadata, 'ocr_text':self.metadata['ocr_text']}
    def _load_audio(self, file_path, **kwargs):
        try:
            import librosa
            audio_data, sr = librosa.load(file_path, **kwargs)
            self.metadata['sample_rate']=sr; self.metadata['duration']=len(audio_data)/sr
            self.metadata['channels']=audio_data.shape[0] if len(audio_data.shape)>1 else 1
            return {'audio_data':audio_data, 'sample_rate':sr, 'metadata':self.metadata}
        except ImportError:
            raise ImportError("librosa required for audio")
    def _load_video(self, file_path, **kwargs):
        try:
            import cv2
            cap = cv2.VideoCapture(str(file_path))
            frames = []
            while True:
                ret, frame = cap.read()
                if not ret: break
                frames.append(frame)
            cap.release()
            self.metadata['total_frames']=len(frames)
            self.metadata['frame_shape']=frames[0].shape if frames else None
            self.metadata['fps']=cap.get(cv2.CAP_PROP_FPS)
            return {'frames':frames, 'metadata':self.metadata}
        except ImportError:
            raise ImportError("OpenCV required for video")
    def _load_ml_model(self, file_path, **kwargs):
        ext = file_path.suffix.lower()
        model_data = None
        if ext in ['.pkl','.pickle']:
            with open(file_path,'rb') as f: model_data = pickle.load(f)
            self.metadata['model_type']='pickle'
        elif ext == '.joblib':
            import joblib; model_data = joblib.load(file_path)
            self.metadata['model_type']='joblib'
        elif ext == '.h5':
            import tensorflow as tf; model_data = tf.keras.models.load_model(file_path)
            self.metadata['model_type']='h5'
        elif ext in ['.pt','.pth']:
            import torch; model_data = torch.load(file_path, map_location='cpu')
            self.metadata['model_type']='pytorch'
        return {'model':model_data, 'type':self.metadata.get('model_type'), 'metadata':self.metadata}
    def _load_notebook(self, file_path):
        import nbformat
        with open(file_path,'r',encoding='utf-8') as f:
            nb = nbformat.read(f, as_version=4)
        cells = [{'type':cell.cell_type, 'source':cell.source} for cell in nb.cells]
        return {'notebook':nb, 'cells':cells, 'metadata':nb.metadata}
    def _load_binary(self, file_path):
        with open(file_path,'rb') as f: data=f.read()
        return {'data':data, 'size':len(data), 'hex_preview':data[:100].hex() if data else ""}
    def _load_generic(self, file_path):
        with open(file_path,'r', errors='ignore') as f: content=f.read()
        return {'content':content, 'type':'generic', 'size':len(content)}